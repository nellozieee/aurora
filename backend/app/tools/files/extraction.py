"""Best-effort text extraction for common document formats.

Used by the read_file tool so the model gets readable text back for PDFs,
DOCX, and XLSX files, not raw bytes it cannot use.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

TEXT_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".html", ".htm", ".py", ".js", ".ts", ".yml", ".yaml"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx"}


class UnsupportedFileTypeError(Exception):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    if suffix == ".csv":
        return _extract_csv(path)
    if suffix in {".html", ".htm"}:
        return _extract_html(path)
    if suffix in TEXT_EXTENSIONS or suffix == "":
        return path.read_text(encoding="utf-8", errors="replace")

    raise UnsupportedFileTypeError(
        f"Unsupported file type '{suffix}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
    )


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            lines.append(", ".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(lines)


def _extract_csv(path: Path) -> str:
    with path.open(encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        return "\n".join(", ".join(row) for row in reader)


def _extract_html(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


def validate_json(text: str) -> None:
    json.loads(text)  # raises json.JSONDecodeError if malformed


def io_text_size(text: str) -> int:
    return len(text.encode("utf-8", errors="replace"))
